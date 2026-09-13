#!/usr/bin/env python3
"""Render the manuscript from reviewed Markdown; reuse the editable Table 1."""
import copy, pathlib, re, subprocess, tempfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE=pathlib.Path(__file__).resolve().parent
text=(HERE/'barbac_manuscript.md').read_text()
text=re.sub(r'!\[Figure ([12])[.] ([^\n]+)\]\(([^)]+)\)\{width=178mm\}',
  lambda m: f'![Workflow figure {m[1]}]({m[3]}){{width=178mm}}\n\n**Figure {m[1]}.** {m[2]}',text)
assert '{{WORKFLOW_RESULT}}' not in text, 'Insert the completed workflow receipt first.'
text=re.sub(r'^\*\*((?:[1-5](?:\.[0-9]+)? [^*]+)|(?:Data and software availability|Acknowledgements|Funding|Competing interests|References))\*\*$',
            lambda m: ('### ' if re.match(r'[1-5]\.',m[1]) else '## ')+m[1], text, flags=re.M)
table_pattern=r'!\[Table 1\.[^\n]+\]\(publication_tables/table_1_benchmark_comparison.png\)\{width=178mm\}'
assert len(re.findall(table_pattern,text))==1
with tempfile.TemporaryDirectory(prefix='barbac-manuscript-') as folder:
    temp=pathlib.Path(folder)
    doc_text=re.sub(table_pattern,'BARBAC_EDITABLE_TABLE_ONE',text)
    (temp/'word.md').write_text(doc_text)
    output=HERE/'barbac_manuscript_publication.docx'
    subprocess.run(['quarto','pandoc',str(temp/'word.md'),'--from=markdown-implicit_figures',
                    '--resource-path='+str(HERE),'-o',str(output)],cwd=HERE,check=True)
    doc=Document(output)
    for sec in doc.sections:
        sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
        sec.left_margin=sec.right_margin=Inches(.62)
        sec.top_margin=sec.bottom_margin=Inches(.62)
    for name in ('Normal','Body Text','First Paragraph'):
        style=doc.styles[name]
        style.font.name='Times New Roman';style.font.size=Pt(11)
        style.paragraph_format.space_after=Pt(5)
    for style in doc.styles:
        if style.name in ('Title','Heading 1','Heading 2','Heading 3'):
            style.font.name='Times New Roman'
            style.font.color.rgb=RGBColor.from_string('173E3B')
    tables=Document(HERE/'publication_tables/table_1_benchmark_comparison.docx').tables
    assert len(tables)==1
    target=[p for p in doc.paragraphs if p.text=='BARBAC_EDITABLE_TABLE_ONE']
    assert len(target)==1
    before=target[0].insert_paragraph_before()
    before.paragraph_format.page_break_before=True
    target[0]._p.addnext(copy.deepcopy(tables[0]._tbl))
    target[0]._p.getparent().remove(target[0]._p)
    for sec in doc.sections:
        para=sec.footer.paragraphs[0];para.alignment=2
        run=para.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
        run._r.addnext(field)
    doc.save(output)
    pdf_text=text.replace('](figures/workflow.png)','](figures/workflow.pdf)').replace('](figures/workflow_application.png)','](figures/workflow_application.pdf)')
    pdf_text=re.sub(table_pattern,lambda _: '\\clearpage\n\n![](publication_tables/table_1_benchmark_comparison.png){width=166mm}',pdf_text)
    pdf_text=re.sub(r'\*\*Table 1[.]\*\*.*?(?=### 4[.]2)', lambda _: '\\clearpage\n\n',pdf_text,flags=re.S)
    (temp/'pdf.md').write_text(pdf_text)
    subprocess.run(['quarto','pandoc',str(temp/'pdf.md'),'--from=markdown-implicit_figures',
      '--resource-path='+str(HERE),'--pdf-engine=xelatex','-V','mainfont=Times New Roman',
      '-V','geometry:margin=20mm','-V','fontsize=11pt','-V','colorlinks=true',
      '-o',str(HERE/'barbac_manuscript_publication.pdf')],cwd=HERE,check=True)
print('Rendered publication DOCX with editable orange Table 1, and manuscript PDF.')
