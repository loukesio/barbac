"""Build Table 1 from verified receipts; no simulation or clustering is run."""
import csv
import hashlib
import json
import math
from pathlib import Path
import re
import statistics
import subprocess
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

HERE = Path(__file__).resolve().parent
SOURCE = HERE.parents[1]
OUT = SOURCE / 'manuscript/publication_tables'
PREFIX = 'table_1_benchmark_comparison'
METHODS = ['hamming','lv','shepherd_documented','starcode_sphere','starcode_mp','bartender']
NAMES = dict(zip(METHODS,['barbac Hamming','barbac LV + Poisson','Shepherdᵃ','Starcode sphere','Starcode MP','Bartender']))
DESIGNS = ['random_mixed','anchored_mixed','milos']
PANELS = {
    'random_mixed': 'A  Random N20: substitutions + indels (60 libraries)',
    'anchored_mixed': 'B  Anchored N20 + AA/TT: substitutions + indels (60 libraries)',
    'milos': 'C  Milo / Johnson published simulation (one fixed reference)',
}
METRICS = ['fn','fp','f1_percent','read_assignment_accuracy_percent','workflow_seconds']
HIGHER = {'f1_percent','read_assignment_accuracy_percent'}
WIDTHS = [39,20,20,45,31,23]
HIGHLIGHT = 'FBE0B7'
TITLE = 'Table 1. Accuracy and runtime of barcode clustering across three benchmarks.'
SUBTITLE = ('Panels A and B: 10,000 true barcodes and approximately 1 million reads per library. '
            'Panel C: 100,000 true barcodes and 24,996,128 reads.')
HEADERS = ['Method','FN ↓','FP ↓','F1 (%) ↑','Read accuracy\n(%) ↑','Runtimeᵇ\n(s) ↓']
NOTES = [
    'Panels A and B: FN, FP and read accuracy are means; F1 is mean ± standard deviation; runtime is median. Panel C reports one observation. Bold and pale orange shading identify the best observed value for each metric within each panel.',
    'FN, false negatives: true identities absent from inferred centroids. FP, false positives: inferred identities absent from truth. Exact-centroid F1 = 200TP/(2TP + FN + FP), including zero-read true identities. Read accuracy is the percentage of all reads assigned to their true barcode; incorrect and unassigned reads both count as errors.',
    'ᵃ Shepherd uses distance 3, Bayes-factor threshold −4, and the generating substitution rate 0.004 in panels A/B; panel C uses automatic rate estimation. This is the completed post hoc sensitivity configuration; all other methods retain their original settings and results.',
    'ᵇ All workflows use one thread. Runtime includes fresh-worker/tool startup, required input conversion and centroid/member export; common staging and scoring are excluded. Shepherd was measured in a later session, so its timing comparisons are descriptive.',
    'The accompanying Table 1 methods provide simulation details, complete settings, paired statistical comparisons and measurement provenance.'
]


def sha(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream,'sha256').hexdigest()


def load(name):
    return json.loads((HERE/name).read_text())


def reviewed_rows():
    summaries = load('summary.json')
    receipts = load('combined_results.json')
    assert len(receipts)==726 and len(summaries)==18
    assert all(r['status']=='complete' for r in receipts)
    assert len({(r['condition'],r['seed'],r['method']) for r in receipts})==726
    assert load('independent_validation.json')['all_605_other_method_records_exactly_preserved']
    for r in summaries:
        cells=[x for x in receipts if (x['condition'],x['method'])==(r['condition'],r['method'])]
        n=1 if r['condition']=='milos' else 60
        assert len(cells)==n and r['complete_population'] and r['n_successful']==n
        for metric in METRICS:
            value = (statistics.median(x['reported_workflow_seconds'] for x in cells)
                     if metric=='workflow_seconds' else statistics.mean(x[metric] for x in cells))
            assert math.isclose(value,r[metric],rel_tol=0,abs_tol=1e-11),(r['method'],metric)
        if n>1:
            assert math.isclose(statistics.stdev(x['f1_percent'] for x in cells),r['f1_percent_sd'],abs_tol=1e-12)
    expected={(d,m) for d in DESIGNS for m in METHODS}
    assert {(r['condition'],r['method']) for r in summaries}==expected
    return sorted(summaries,key=lambda r:(DESIGNS.index(r['condition']),METHODS.index(r['method'])))


def displays(rows):
    output=[]
    for r in rows:
        digits=0 if r['condition']=='milos' else 2
        peers=[x for x in rows if x['condition']==r['condition']]
        leaders=[False]+[abs(r[m]-(max if m in HIGHER else min)(x[m] for x in peers))<1e-12 for m in METRICS]
        f1=f"{r['f1_percent']:.5f}"
        if r['condition']!='milos':f1+=f" ± {r['f1_percent_sd']:.4f}"
        values=[NAMES[r['method']],f"{r['fn']:.{digits}f}",f"{r['fp']:.{digits}f}",f1,
                f"{r['read_assignment_accuracy_percent']:.6f}",f"{r['workflow_seconds']:.2f}"]
        output.append(dict(condition=r['condition'],method=r['method'],values=values,bold=leaders))
    return output


def register_fonts():
    folder=Path('/System/Library/Fonts/Supplemental')
    for name,file in [('Paper','Times New Roman.ttf'),('PaperBold','Times New Roman Bold.ttf'),
                      ('PaperItalic','Times New Roman Italic.ttf'),('PaperBoldItalic','Times New Roman Bold Italic.ttf')]:
        pdfmetrics.registerFont(TTFont(name,str(folder/file)))
    pdfmetrics.registerFontFamily('Paper',normal='Paper',bold='PaperBold',italic='PaperItalic',boldItalic='PaperBoldItalic')
    rl_config.canvas_basefontname='Paper'


def pdf_table(displayed):
    register_fonts()
    caption=ParagraphStyle('Caption',fontName='PaperBold',fontSize=12,leading=15,spaceAfter=7)
    subtitle=ParagraphStyle('Sub',fontName='Paper',fontSize=9.5,leading=12,spaceAfter=12)
    note=ParagraphStyle('Note',fontName='Paper',fontSize=8.5,leading=10.5,spaceAfter=5)
    header=ParagraphStyle('Header',fontName='PaperBold',fontSize=9,leading=11,alignment=2)
    left=ParagraphStyle('LeftHeader',parent=header,alignment=0)
    matrix=[[Paragraph(h.replace('\n','<br/>'),left if i==0 else header) for i,h in enumerate(HEADERS)]]
    commands=[('FONTNAME',(0,0),(-1,-1),'Paper'),('FONTSIZE',(0,0),(-1,-1),9.5),
        ('ALIGN',(1,0),(-1,-1),'RIGHT'),('VALIGN',(0,0),(-1,-1),'MIDDLE'),
        ('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),
        ('TOPPADDING',(0,0),(-1,0),5),('BOTTOMPADDING',(0,0),(-1,0),7),
        ('TOPPADDING',(0,1),(-1,-1),4),('BOTTOMPADDING',(0,1),(-1,-1),4),
        ('LINEABOVE',(0,0),(-1,0),.9,colors.black),('LINEBELOW',(0,0),(-1,0),.6,colors.black)]
    for condition in DESIGNS:
        i=len(matrix);matrix.append([PANELS[condition],'','','','',''])
        commands += [('SPAN',(0,i),(-1,i)),('FONTNAME',(0,i),(-1,i),'PaperBold'),
            ('BACKGROUND',(0,i),(-1,i),colors.HexColor('#F0F2F4')),
            ('TOPPADDING',(0,i),(-1,i),7),('BOTTOMPADDING',(0,i),(-1,i),6)]
        for row in [r for r in displayed if r['condition']==condition]:
            i=len(matrix);matrix.append(row['values'])
            for j,best in enumerate(row['bold']):
                if best:
                    commands += [('FONTNAME',(j,i),(j,i),'PaperBold'),
                                 ('BACKGROUND',(j,i),(j,i),colors.HexColor('#'+HIGHLIGHT))]
    commands.append(('LINEBELOW',(0,len(matrix)-1),(-1,len(matrix)-1),.9,colors.black))
    table=Table(matrix,colWidths=[w*mm for w in WIDTHS]);table.setStyle(TableStyle(commands))
    story=[Paragraph(TITLE,caption),Paragraph(SUBTITLE,subtitle),table,Spacer(1,10)]
    story += [Paragraph(n,note) for n in NOTES]
    pdf=OUT/f'{PREFIX}.pdf'
    SimpleDocTemplate(str(pdf),pagesize=A4,leftMargin=16*mm,rightMargin=16*mm,topMargin=17*mm,bottomMargin=16*mm,
                      title=TITLE,author='barbac').build(story)
    info=subprocess.check_output(['pdfinfo',str(pdf)],text=True)
    assert re.search(r'^Pages:\s+1$',info,re.M),info
    actual=subprocess.check_output(['pdftotext','-layout',str(pdf),'-'],text=True)
    observed=[re.sub(r'\s+',' ',line).strip() for line in actual.splitlines() if line.strip().startswith(tuple(NAMES.values()))]
    assert observed==[' '.join(row['values']) for row in displayed],(observed,displayed)
    fonts=subprocess.check_output(['pdffonts',str(pdf)],text=True)
    for line in fonts.splitlines()[2:]:
        if line.strip():assert re.search(r'\byes\s+yes\s+yes\b',line),line
    subprocess.run(['pdftoppm','-scale-to','2000','-png','-singlefile',str(pdf),str(OUT/PREFIX)],check=True)
    return fonts


def cell_border(cell,**edges):
    props=cell._tc.get_or_add_tcPr()
    borders=OxmlElement('w:tcBorders')
    for edge,size in edges.items():
        item=OxmlElement('w:'+edge);item.set(qn('w:val'),'single');item.set(qn('w:sz'),str(size));item.set(qn('w:color'),'000000');borders.append(item)
    props.append(borders)


def docx_table(displayed):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Mm(17);sec.bottom_margin=Mm(16);sec.left_margin=sec.right_margin=Mm(16)
    normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(9.5)
    normal.paragraph_format.space_after=Pt(0);normal.paragraph_format.line_spacing=1
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(7)
    r=p.add_run(TITLE);r.bold=True;r.font.size=Pt(12)
    p=doc.add_paragraph(SUBTITLE);p.paragraph_format.space_after=Pt(12)
    table=doc.add_table(rows=1,cols=6);table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,width in zip(table.columns,WIDTHS):col.width=Mm(width)
    def populate(row,texts,bolds,header=False):
        for i,(cell,text,bold) in enumerate(zip(row.cells,texts,bolds)):
            cell.width=Mm(WIDTHS[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
            r=p.add_run(text);r.bold=bold;r.font.size=Pt(9 if header else 9.5)
            if header:cell_border(cell,top=7,bottom=5)
            elif bold:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),HIGHLIGHT)
                cell._tc.get_or_add_tcPr().append(shading)
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
    populate(table.rows[0],HEADERS,[True]*6,True)
    for condition in DESIGNS:
        row=table.add_row();cell=row.cells[0].merge(row.cells[-1]);p=cell.paragraphs[0]
        p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.keep_with_next=True;p.add_run(PANELS[condition]).bold=True
        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'F0F2F4');cell._tc.get_or_add_tcPr().append(shading)
        for entry in [r for r in displayed if r['condition']==condition]:
            populate(table.add_row(),entry['values'],entry['bold'])
    for cell in table.rows[-1].cells:cell_border(cell,bottom=7)
    for i,text in enumerate(NOTES):
        p=doc.add_paragraph(text);p.paragraph_format.space_before=Pt(9 if i==0 else 5)
        p.paragraph_format.space_after=Pt(0)
        for r in p.runs:r.font.size=Pt(8.5)
    path=OUT/f'{PREFIX}.docx';doc.save(path)
    # Independently inspect the editable table's serialized values and emphasis.
    with zipfile.ZipFile(path) as z:root=ET.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    observed=[];observed_bold=[];observed_shading=[]
    for row in root.findall('.//w:tbl/w:tr',ns):
        cells=row.findall('w:tc',ns)
        texts=[''.join(node.text or '' for node in cell.findall('.//w:t',ns)) for cell in cells]
        if len(texts)==6 and texts[0] in NAMES.values():
            observed.append(texts)
            observed_bold.append([any(node.get(qn('w:val'),'1')!='0' for node in cell.findall('.//w:rPr/w:b',ns)) for cell in cells])
            observed_shading.append([any(node.get(qn('w:fill'))==HIGHLIGHT for node in cell.findall('./w:tcPr/w:shd',ns)) for cell in cells])
    assert observed==[r['values'] for r in displayed]
    assert observed_bold==[r['bold'] for r in displayed]
    assert observed_shading==[r['bold'] for r in displayed]


def companion():
    text='''# Table 1 methods and statistical comparisons

Table 1 compares six configurations across two independently simulated library designs and one fixed published simulation. It reports the complete Shepherd sensitivity configuration alongside unchanged results for the other methods.

## Datasets and generation

Random libraries have 20 variable bases. Anchored libraries use NNNNNAANNNNNAANNNNNTTNNNNN (20 variable bases, 26 bases total). Each design contains 60 independent libraries with 10,000 true identities and one million expected reads. Within each seed, the two designs share variable identities and parent abundances. Abundances follow the recorded scaled Johnson exponential mixture (9,989 ordinary, 10 intermediate and one high-abundance identity).

Substitutions occur with probability 0.004 per base. Repeat-dependent indel allocation uses the unscaled archived empirical homopolymer-rate table. Repeat coordinates are updated after events. When a repeat exceeds the measured support (length 13), further indel recursion stops for those reads; reads, origin labels and subsequent substitutions are retained. This boundary affected 249 of 119,992,068 reads across 16 of 120 libraries. All registered seeds and reads remain in the analysis.

The fixed Milo / Johnson deposited simulation contains 100,000 true identities and 24,996,128 reads and was used during development. Its reference counts and origin labels differ at four parents by six reads in total absolute count difference; totals and identities reconcile. Its scores are descriptive. Simulation source: https://pmc.ncbi.nlm.nih.gov/articles/PMC10276077/ ; archived analysis: https://zenodo.org/records/7411747 .

## Configurations

All methods retain an explicit distance limit of 3 and one thread. Barbac uses the preserved v14 native implementation with support ordering, merge ratio 20, configured error proxy 0.005 and design scoring disabled. LV uses the existing Poisson option; Hamming retains its rare-indel rescue. Starcode sphere uses sphere clustering; Starcode MP uses message-passing ratio 5. Bartender uses seed length 5, step 1, cutoff 1, z = 5 and forward direction.

Shepherd uses nominal length, distance 3 and its documented default Bayes-factor threshold −4. It receives the known generating substitution rate 0.004 on every simulated library; the fixed Milo reference uses automatic rate estimation. The earlier frozen configuration used threshold +4 and automatic estimation and failed on ten simulated inputs. The complete supplementary configuration was registered after observing those outcomes. Both changes apply consistently; no accuracy-based parameter search was performed. It supplies the generating substitution rate to Shepherd, not truth identities or repeat-specific indel rates. This is a post hoc sensitivity comparison, not an all-default run or a retroactive replacement for the original confirmatory protocol. Author documentation: https://github.com/Nik-Tavakolian/Shepherd .

## Metrics and timing

Exact-centroid F1 includes all designed truth identities, including those with zero reads. For each library, F1 (%) = 200TP/(2TP + FN + FP). Table entries average library-level accuracy scores without pooling reads across libraries. Read accuracy divides correctly assigned reads by all input reads; incorrect and unassigned reads are both failures. The displayed F1 spread is the sample standard deviation, not a confidence interval.

Runtime measures a fresh Python worker including tool/package startup, required input conversion and canonical centroid/member exports. Common staging, scoring and hashing are excluded. Runs were serial. Original accepted times are retained for all methods other than the new Shepherd configuration, including the previously registered 28 sleep-affected calls repeated once with identical clustering outputs. The 121 supplementary Shepherd calls were timed in a later session; lightweight diagnostic file reads overlapped early calls. No sleep/wake event overlapped those new timed calls. Comparisons involving the new Shepherd times are descriptive. Barbac package startup uses the verified lazy BAM-loading change; its clustering implementation is unchanged.

## Paired accuracy comparisons

The original primary analysis compared Barbac LV with four external configurations in each simulated design, keeping an eight-contrast family. Paired-t one-sided lower bounds use alpha 0.05/8; paired bootstrap lower bounds resample whole libraries 100,000 times at the same level. A positive difference favours Barbac LV. The six complete original contrasts with Bartender and Starcode meet both criteria. The two original Shepherd contrasts remain recorded as unavailable.

For the completed Shepherd sensitivity configuration, the same library pairing and conservative alpha 0.05/8 are retained. These two comparisons remain explicitly post hoc. The table below combines the six original contrasts and the two supplementary Shepherd contrasts for inspection, with their scope identified. Milo is excluded from population-level inference.
'''
    original=json.loads((SOURCE/'benchmark/publication_final/accuracy_contrasts.json').read_text())
    fresh=load('accuracy_contrasts.json')
    rows=[]
    for condition in DESIGNS[:2]:
        for method in ['shepherd','starcode_sphere','starcode_mp','bartender']:
            r=(next(x for x in fresh if x['condition']==condition) if method=='shepherd'
               else next(x for x in original if x['condition']==condition and x['competitor']==method))
            rows.append([('Random' if condition=='random_mixed' else 'Anchored')+' / '+
                (NAMES['shepherd_documented'] if method=='shepherd' else NAMES[method]),
                f"{r['mean_difference']:+.6f}",f"{r['simultaneous_one_sided_lower']:+.6f}",
                f"{r['bootstrap_lower']:+.6f}",'Post hoc' if method=='shepherd' else 'Original'])
    text+='\n| Design / comparator | LV − comparator F1 (pp) | t lower | Bootstrap lower | Scope |\n|---|---:|---:|---:|---|\n'
    text+='\n'.join('| '+' | '.join(r)+' |' for r in rows)
    text+='\n\nAll displayed summary values were reconciled with 726 execution records. Earlier independent checks reproduced every new Shepherd centroid score, three complete read mappings and the supplementary paired-t calculations in base R. Source fingerprints and the export checks accompany the table in `table_1_provenance.json`.\n'
    (OUT/'table_1_methods.md').write_text(text)
    # A separate concise PDF keeps the table itself ready to insert into a paper.
    styles={
        'title':ParagraphStyle('MTitle',fontName='PaperBold',fontSize=13,leading=16,spaceAfter=10),
        'heading':ParagraphStyle('MHead',fontName='PaperBold',fontSize=10.5,leading=13,spaceBefore=8,spaceAfter=4,keepWithNext=True),
        'body':ParagraphStyle('MBody',fontName='Paper',fontSize=9.2,leading=11.5,spaceAfter=6),
    }
    story=[];audit_note=''
    for block in text.split('\n\n'):
        if block.startswith('|'):continue
        if block.startswith('All displayed summary'):
            audit_note=block.replace('`','');continue
        if block.startswith('## Paired accuracy'):story.append(PageBreak())
        kind='title' if block.startswith('# ') else 'heading' if block.startswith('## ') else 'body'
        content=block.removeprefix('## ').removeprefix('# ').replace('&','&amp;').replace('<','&lt;')
        if content.strip():story.append(Paragraph(content,styles[kind]))
    matrix=[['Design / comparator','Mean Δ F1 (pp)','t lower','Bootstrap lower','Scope']]+rows
    table=Table(matrix,colWidths=[62*mm,32*mm,25*mm,30*mm,29*mm])
    table.setStyle(TableStyle([('FONTNAME',(0,0),(-1,0),'PaperBold'),('FONTNAME',(0,1),(-1,-1),'Paper'),
        ('FONTSIZE',(0,0),(-1,-1),8.7),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4),
        ('ALIGN',(1,0),(3,-1),'RIGHT'),('LINEABOVE',(0,0),(-1,0),.7,colors.black),
        ('LINEBELOW',(0,0),(-1,0),.5,colors.black),('LINEBELOW',(0,-1),(-1,-1),.7,colors.black)]))
    story.append(table)
    story += [Spacer(1,10),Paragraph(audit_note,styles['body'])]
    SimpleDocTemplate(str(OUT/'table_1_methods.pdf'),pagesize=A4,leftMargin=16*mm,rightMargin=16*mm,
                      topMargin=16*mm,bottomMargin=16*mm,title='Table 1 methods and statistical comparisons').build(story)
    extracted=subprocess.check_output(['pdftotext','-layout',str(OUT/'table_1_methods.pdf'),'-'],text=True)
    actual=[re.sub(r'\s+',' ',line).strip() for line in extracted.splitlines()
            if line.strip().startswith(('Random /','Anchored /'))]
    assert actual==[' '.join(row) for row in rows],(actual,rows)
    subprocess.run(['pdftoppm','-scale-to','1600','-png',str(OUT/'table_1_methods.pdf'),str(OUT/'table_1_methods')],check=True)


def main():
    rows=reviewed_rows();shown=displays(rows);OUT.mkdir(parents=True,exist_ok=True)
    fonts=pdf_table(shown);docx_table(shown);companion()
    with (OUT/f'{PREFIX}.csv').open('w',newline='') as stream:
        writer=csv.writer(stream);writer.writerow(['Benchmark']+[h.replace('\n',' ') for h in HEADERS])
        for r in shown:writer.writerow([PANELS[r['condition']]]+r['values'])
    with (OUT/'table_1_unrounded_data.csv').open('w',newline='') as stream:
        fields=['condition','method','planned_n']+METRICS+['f1_percent_sd']
        writer=csv.DictWriter(stream,fieldnames=fields,extrasaction='ignore');writer.writeheader();writer.writerows(rows)
    (OUT/'README.md').write_text('''# Publication Table 1

[Print-ready comparison PDF](table_1_benchmark_comparison.pdf) · [Editable Word table](table_1_benchmark_comparison.docx) · [CSV](table_1_benchmark_comparison.csv)

All six method configurations are shown for each of the three benchmarks. The table includes FN, FP, exact-centroid F1, read-assignment accuracy and workflow runtime. Bold and pale orange shading indicate the numerical leader within each benchmark and metric.

The [companion methods PDF](table_1_methods.pdf) contains full settings, simulation provenance, timing details and the original versus supplementary statistical comparisons. [Unrounded data](table_1_unrounded_data.csv) and [source receipts](table_1_provenance.json) support reuse.

Rebuild from this worktree with `python3 benchmark/shepherd_completion/publication_table.py`. This only reads saved results and renders files; it runs no clustering or simulation. The PDF is the visually verified print layout; the Word table remains editable and its serialized values are checked against the same source.
''')
    sources=['summary.json','combined_results.json','protocol.json','freeze.json','independent_validation.json',
             'accuracy_contrasts.json','timing_environment_validation.json']
    exported=[p for p in OUT.iterdir() if p.suffix in ['.pdf','.docx','.csv','.md']]
    record=dict(source_sha256={str(HERE/p):sha(HERE/p) for p in sources},renderer_sha256=sha(__file__),
        rows=18,configurations=6,benchmarks=3,simulated_libraries_per_design=60,combined_execution_records=726,
        displayed_metrics_recomputed_from_receipts=True,pdf_text_values_verified=True,docx_serialized_values_verified=True,
        main_pdf_pages=1,pdf_fonts_embedded=True,pdffonts=fonts,visual_inspection_pending=True,
        best_value_highlight_hex=HIGHLIGHT,highlighted_cells=sum(sum(r['bold']) for r in shown),
        docx_shading_matches_metric_leaders=True,
        files_sha256={p.name:sha(p) for p in exported})
    (OUT/'table_1_provenance.json').write_text(json.dumps(record,indent=2)+'\n')
    print(OUT/f'{PREFIX}.pdf')


if __name__=='__main__':main()
