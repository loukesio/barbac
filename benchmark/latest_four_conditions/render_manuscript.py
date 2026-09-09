"""Regenerate the editable Word manuscript and verify every benchmark cell."""
import json
import hashlib
from pathlib import Path
import re
import subprocess
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt

from run_comparison import HERE, ROOT, sha


def main():
    source=ROOT/'manuscript/barbac_manuscript.md'
    reference=ROOT/'manuscript/barbac_manuscript.docx'
    output=ROOT/'manuscript/barbac_manuscript_updated.docx'
    subprocess.run(['pandoc',str(source),'--from','markdown','--to','docx',
                    '--resource-path',str(source.parent),'--reference-doc',str(reference),
                    '--output',str(output)],check=True)
    # Pandoc preserves the reference's embedded .ttf parts but omits their
    # content type. Restore only missing declarations from the base document.
    with ZipFile(reference) as archive:
        base_types=ET.fromstring(archive.read('[Content_Types].xml'))
    with ZipFile(output) as archive:
        entries=[(info,archive.read(info.filename)) for info in archive.infolist()]
    types=ET.fromstring(dict((info.filename,data) for info,data in entries)['[Content_Types].xml'])
    extensions={el.attrib['Extension'] for el in types if el.tag.endswith('Default')}
    present={info.filename.rsplit('.',1)[-1] for info,data in entries if '.' in info.filename}
    restored=[]
    for el in base_types:
        extension=el.attrib.get('Extension')
        if extension and extension in present and extension not in extensions:
            types.append(el);restored.append(extension)
    with ZipFile(output,'w') as archive:
        for info,data in entries:
            archive.writestr(info,ET.tostring(types,encoding='utf-8',xml_declaration=True)
                             if info.filename=='[Content_Types].xml' else data)
    document=Document(output)
    expected=[]
    for line in (HERE/'paper_table.md').read_text().splitlines():
        if line.startswith('|') and not line.startswith('|---'):
            expected.append([cell.strip().replace('**','') for cell in line.strip('|').split('|')])
    tables=[t for t in document.tables if [c.text for c in t.rows[0].cells]==expected[0]]
    assert len(tables)==1
    table=tables[0]
    assert [[c.text for c in row.cells] for row in table.rows]==expected
    table.autofit=False
    widths=[0.60,1.85,0.75,1.0,1.05,0.95]
    for column,width in zip(table.columns,widths): column.width=Inches(width)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell,width in zip(row.cells,widths):
            cell.width=Inches(width)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2)
                paragraph.paragraph_format.space_before=Pt(2)
                for run in paragraph.runs: run.font.size=Pt(9)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    series_expected=[]
    if 'media/time_series_trajectories.png' in source.read_text():
        series_section=ROOT/'benchmark/time_series_chen2023/results/paper_time_series_section.md'
        for line in series_section.read_text().splitlines():
            if line.startswith('|') and not line.startswith('|---'):
                series_expected.append([cell.strip() for cell in line.strip('|').split('|')])
        series_table=next(t for t in document.tables if [c.text for c in t.rows[0].cells]==series_expected[0])
        assert [[c.text for c in row.cells] for row in series_table.rows]==series_expected
        series_table.autofit=False
        series_widths=[1.8,0.95,1.35,1.05,1.05]
        for column,width in zip(series_table.columns,series_widths):column.width=Inches(width)
        for row in series_table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,width in zip(row.cells,series_widths):
                cell.width=Inches(width)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after=Pt(2)
                    for run in paragraph.runs:run.font.size=Pt(9)
        series_table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    document.save(output)
    reopened=Document(output)
    actual=next(t for t in reopened.tables if [c.text for c in t.rows[0].cells]==expected[0])
    assert [[c.text for c in row.cells] for row in actual.rows]==expected
    if series_expected:
        actual_series=next(t for t in reopened.tables if [c.text for c in t.rows[0].cells]==series_expected[0])
        assert [[c.text for c in row.cells] for row in actual_series.rows]==series_expected
    # python-docx's inline_shapes collection omits pictures inside hyperlinks;
    # inspect every actual image reference, including the equation images.
    image_hashes=[]
    for blip in reopened.element.xpath('.//a:blip'):
        rid=blip.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        image_hashes.append(hashlib.sha256(reopened.part.rels[rid].target_part.blob).hexdigest())
    source_hashes=[sha(source.parent/path) for path in re.findall(r'!\[\]\(([^)]+)\)',source.read_text())]
    expected_images=9 if 'media/time_series_trajectories.png' in source.read_text() else 8
    assert sorted(image_hashes)==sorted(source_hashes) and len(image_hashes)==expected_images
    receipt=dict(status='passed',table_data_rows=len(expected)-1,table_columns=len(expected[0]),
                 all_cells_match_markdown=True,embedded_figures=len(image_hashes),
                 restored_reference_content_types=restored,
                 time_series_table_data_cells=(len(series_expected)-1)*len(series_expected[0]) if series_expected else 0,
                 source_sha256=sha(source),output_sha256=sha(output),
                 validation='DOCX opens successfully; all 180 data cells match the paper table; repeated headers and fixed column widths applied. All source images, including benchmark/time-series figures when present, are embedded with matching hashes. No visual Word/PDF rendering performed.')
    (HERE/'manuscript_validation.json').write_text(json.dumps(receipt,indent=2)+'\n')
    print(json.dumps(receipt,indent=2))


if __name__=='__main__': main()
